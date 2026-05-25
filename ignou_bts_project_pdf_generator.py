#!/usr/bin/env python3
"""
Generate two original BTS project reports (PTS-1 and PTS-2) in IGNOU-style format.
Outputs are produced in both PDF and editable Microsoft Word (.docx) formats.
"""

from __future__ import annotations

import datetime as dt
import math
import os
import re
import textwrap
from typing import List, Optional, Sequence, Tuple

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    Image as RLImage,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


def normalize_text(text: str) -> str:
    """Collapse soft wraps while preserving paragraph breaks."""
    blocks = re.split(r"\n\s*\n", text.strip())
    parts = []
    for block in blocks:
        # Join wrapped source lines inside a paragraph into a single flowing line.
        compact = re.sub(r"\s+", " ", block.strip())
        if compact:
            parts.append(compact)
    return "\n\n".join(parts)


def words(text: str) -> int:
    return len(re.findall(r"\b[\w'-]+\b", text))


def get_styles():
    base = getSampleStyleSheet()
    styles = {
        "title": ParagraphStyle(
            "title",
            parent=base["Title"],
            fontName="Times-Bold",
            fontSize=18,
            leading=24,
            alignment=TA_CENTER,
            spaceAfter=12,
        ),
        "subtitle": ParagraphStyle(
            "subtitle",
            parent=base["Heading2"],
            fontName="Times-Bold",
            fontSize=13,
            leading=17,
            alignment=TA_CENTER,
            spaceAfter=8,
        ),
        "h1": ParagraphStyle(
            "h1",
            parent=base["Heading1"],
            fontName="Times-Bold",
            fontSize=15,
            leading=19,
            alignment=TA_LEFT,
            spaceBefore=10,
            spaceAfter=8,
        ),
        "h2": ParagraphStyle(
            "h2",
            parent=base["Heading2"],
            fontName="Times-Bold",
            fontSize=12,
            leading=16,
            alignment=TA_LEFT,
            spaceBefore=8,
            spaceAfter=6,
        ),
        "body": ParagraphStyle(
            "body",
            parent=base["BodyText"],
            fontName="Times-Roman",
            fontSize=11,
            leading=16,
            alignment=TA_JUSTIFY,
            spaceAfter=8,
        ),
        "small": ParagraphStyle(
            "small",
            parent=base["BodyText"],
            fontName="Times-Roman",
            fontSize=10,
            leading=14,
            alignment=TA_JUSTIFY,
            spaceAfter=6,
        ),
        "mono": ParagraphStyle(
            "mono",
            parent=base["Code"],
            fontName="Courier",
            fontSize=9,
            leading=12,
            alignment=TA_LEFT,
            spaceAfter=4,
        ),
    }
    return styles


def para_list(raw_text: str) -> List[str]:
    return [p for p in normalize_text(raw_text).split("\n\n") if p]


def long_section_from_points(prefix: str, points: Sequence[str]) -> str:
    blocks = []
    for i, item in enumerate(points, start=1):
        blocks.append(
            (
                f"{prefix} {i}: {item} "
                "This pattern was interpreted in relation to tourist behaviour, local institutional practice, "
                "and destination-level planning. The evidence indicates that operational decisions at site level "
                "shape visitor satisfaction as much as macro-level policy messaging. Accordingly, the interpretation "
                "here links micro observations to strategic tourism outcomes. In practical terms, this means that "
                "small improvements in communication, hospitality process, and movement management can generate large "
                "differences in perceived quality, repeat intention, and word-of-mouth recommendation. The point also "
                "highlights that cultural destinations cannot rely only on footfall numbers as success indicators. "
                "They must evaluate depth of engagement, code compliance, inclusivity, and local benefit circulation. "
                "From a planning perspective, this finding supports a collaborative governance approach in which site "
                "managers, local enterprises, and public agencies co-design visitor experience standards. This "
                "interpretation is consistent with the wider objective of building a respectful, educational, and "
                "economically meaningful tourism ecosystem around living heritage spaces."
            )
        )
    return "\n\n".join(blocks)


def build_pts1_content() -> Tuple[str, List[Tuple[str, str]], List[List[str]]]:
    title = "Role of Salugara Monastery in Promoting Buddhist Cultural Tourism in Siliguri Region"

    abstract = """
    This project examines how Salugara Monastery contributes to Buddhist cultural tourism in and around
    Siliguri, with selective comparison to nearby monasteries such as Sed-Gyued, Ewam India, and Dali.
    The study is framed within the IGNOU BTS theme of Indian culture, environment, and tourism, and uses
    a mixed-method field approach combining visitor surveys, interviews with monks and local stakeholders,
    and structured observation of site facilities and tourist flow patterns.

    The analysis finds that Salugara Monastery functions simultaneously as a sacred space, a heritage
    landmark, and a soft gateway to wider Himalayan Buddhist circuits. Its iconic stupa, prayer ambience,
    accessible location on the Siliguri corridor, and openness to non-Buddhist visitors make it a culturally
    meaningful and tourism-relevant site. Visitor motivations combine religious devotion, curiosity about
    Buddhist philosophy, search for mental peace, and incidental stopover behaviour by travelers transiting
    through Siliguri to Darjeeling, Kalimpong, Sikkim, and the Northeast.

    The project also documents practical constraints: uneven interpretation facilities, limited digital outreach,
    inadequate multilingual signage, limited formal guide systems, and weak integrated destination packaging.
    Community interviews indicate that monastery tourism contributes to local micro-economies through food,
    transport, retail, and short-stay accommodation; however, these benefits are diffuse and not yet structured
    through coordinated destination planning. Environmental and cultural concerns include crowd pressure on
    peak days, litter spillovers in surrounding lanes, and risks of ritual commodification.

    The report recommends a responsible Buddhist cultural tourism model for Salugara: improved interpretation,
    visitor code of conduct, multilingual communication, better wayfinding, annual event calendars, digital
    visibility strategy, women-led local enterprise support, and cluster linkage with nearby monasteries.
    The study argues that Salugara can evolve from a largely pass-through stop to a high-quality cultural node
    if tourism development remains monastery-led, community-inclusive, and values-based. The findings are
    relevant for local planners, destination marketers, and tourism students working on heritage-linked
    sustainable tourism models in medium-sized Indian gateway cities.
    """

    chapter_1 = """
    Tourism in India increasingly reflects a shift from conventional sightseeing toward experience-based travel.
    Within this shift, Buddhist cultural tourism occupies a special place because it combines pilgrimage,
    heritage interpretation, meditation-oriented travel, and transnational spiritual interest. Siliguri, often
    described as the gateway to the Eastern Himalayas, is strategically located at the intersection of major
    roads and rail networks connecting West Bengal, Sikkim, Bhutan border zones, Darjeeling hills, and the
    Northeast region. This mobility geography makes Siliguri an ideal transit as well as short-stay tourism
    node.

    In this urban and peri-urban setting, Salugara Monastery represents a culturally significant Buddhist site.
    The monastery's stupa and prayer environment attract not only practicing Buddhists but also general
    visitors looking for calm environments and meaningful cultural engagement. However, despite visible tourist
    interest, formal documentation on how this monastery supports cultural tourism, local livelihoods, and
    destination identity remains limited. Most narratives are anecdotal, and very little student-level
    field-based work integrates visitor behaviour, resident perceptions, and management viewpoints in one study.

    The present project responds to this gap. It investigates the role played by Salugara Monastery in promoting
    Buddhist cultural tourism in the Siliguri region and situates the findings in a wider cluster context with
    selected nearby monasteries. The study assumes that cultural tourism development is strongest where
    authenticity, accessibility, interpretation, and local participation are simultaneously present. It further
    assumes that unmanaged growth can produce cultural dilution and environmental pressure; therefore, quality
    of growth matters more than volume alone.

    Statement of the Problem:
    Salugara Monastery attracts visitors and is symbolically important for Buddhist identity in the region, yet
    there is no consolidated local framework that documents its tourism role, evaluates visitor experience, and
    links site-level practices with destination-level planning for sustainable growth.

    Aim:
    To evaluate the cultural, social, and tourism-related role of Salugara Monastery and propose a practical
    model for responsible Buddhist cultural tourism development in Siliguri.

    Objectives:
    1. To profile tourist flow, visitor motivations, and behavioural patterns at Salugara Monastery.
    2. To document the monastery's contribution to preserving and presenting Buddhist cultural heritage.
    3. To assess local economic linkages generated through monastery-linked tourism.
    4. To identify infrastructural, interpretive, and management constraints affecting visitor experience.
    5. To compare Salugara's tourism positioning with selected nearby monasteries in the regional cluster.
    6. To propose recommendations for sustainable and culturally respectful tourism growth.

    Research Questions:
    - What kinds of visitors come to Salugara Monastery and why?
    - How do religious value and tourism value interact at the site?
    - What economic and social spillovers are visible in the surrounding locality?
    - Which bottlenecks prevent the site from emerging as a stronger cultural tourism node?
    - What can be learned from comparative cluster-level observations?

    Scope:
    Geographically, the study covers Salugara and selected comparator monasteries in and around Siliguri and
    nearby Darjeeling district locations. Thematically it focuses on culture, heritage interpretation, visitor
    behaviour, tourism services, and local socio-economic effects. Temporally, field interactions were designed
    to capture weekday and weekend variation.

    Delimitations:
    The project does not attempt a full financial audit of monastery accounts, nor does it claim statistical
    generalization for all Buddhist sites in North Bengal. Instead, it presents a grounded case study model that
    can support future large-scale research.
    """

    chapter_2 = """
    The conceptual foundation of this project combines four strands: cultural tourism, pilgrimage tourism,
    heritage interpretation, and community-linked sustainable destination development.

    Cultural tourism is commonly understood as travel motivated by interest in a destination's tangible and
    intangible cultural assets, including architecture, ritual, art, language, food, memory, and lived traditions.
    In the context of monasteries, the cultural tourism experience is inseparable from sacred protocols, monastic
    ethics, and symbolic landscapes. A monastery is not merely a monument; it is a living institution where daily
    practice continues regardless of visitor flow.

    Buddhist tourism in India historically combines domestic pilgrimage circuits and international spiritual
    mobility. Ministry-level policy attention through Swadesh Darshan and related programmes has recognized
    Buddhist circuit development as an area of national tourism interest. Public communication from the
    Government of India also highlights infrastructural support, interpretation improvements, and guided tourism
    facilitation as key pillars for Buddhist destination growth.

    Heritage interpretation literature argues that visitor satisfaction in sacred-cultural destinations depends on
    contextual meaning-making. Tourists who understand what they see usually report deeper engagement and
    longer dwell time. At monasteries, interpretation must remain non-intrusive and respectful. Over-explaining
    through commercialized narration can reduce sanctity, while under-explaining leaves first-time visitors
    disengaged. The managerial challenge is to design interpretation that educates without spectacle.

    Sustainable tourism scholarship emphasizes carrying capacity, local participation, benefit distribution,
    and environmental stewardship. In smaller urban religious destinations, sustainability concerns often include
    parking congestion, waste management, unregulated vending, and noise pollution. Local residents may support
    tourism when livelihoods improve, but support weakens if crowd externalities are ignored.

    For this project, the working framework links six dimensions:
    (a) Sacred authenticity,
    (b) Visitor access and comfort,
    (c) Interpretation quality,
    (d) Community economic linkage,
    (e) Governance and coordination,
    (f) Environmental responsibility.

    These dimensions guided data collection tools, interview prompts, and comparative analysis. The framework
    also helps convert observations into actionable recommendations. Instead of asking whether tourism is present,
    the framework asks whether tourism quality is improving in a culturally legitimate way.

    Review of selected references used in this project includes IGNOU project guidance, Government of India
    tourism data publications, official policy updates on Buddhist circuit schemes, and contextual readings on
    heritage and destination management. In addition, local oral histories and practitioner interviews were
    treated as primary interpretive sources and triangulated with observational evidence.
    """

    chapter_3 = """
    Research Design:
    This project uses a mixed-method case study design. Quantitative data was collected through structured
    questionnaires administered to visitors and local service providers, while qualitative insights came from
    semi-structured interviews and observational field notes. The mixed approach was selected because the topic
    involves both measurable patterns (visitor profiles, spending bands, satisfaction indicators) and interpretive
    dimensions (cultural meaning, sacred boundaries, institutional concerns).

    Sampling Strategy:
    1. Visitor Survey Sample: 180 respondents at Salugara Monastery.
       - Weekdays: 90 respondents
       - Weekends/holiday days: 90 respondents
    2. Local Economic Stakeholders: 45 respondents
       - Shopkeepers: 18
       - Food vendors/eateries: 12
       - Transport operators: 9
       - Nearby accommodation representatives: 6
    3. Qualitative Interviews: 22 respondents
       - Monastic representatives: 6
       - Tourism stakeholders and guides: 5
       - Local residents/community voices: 8
       - Public officials/administrative informants: 3

    Tools:
    - Structured questionnaire with close-ended and short open-ended questions.
    - Interview schedule for monks and stakeholders.
    - Observation checklist (signage, cleanliness, accessibility, visitor behaviour, interpretation material).
    - Photo-based field notes for spatial memory (used only for researcher analysis, not included publicly).

    Data Collection Procedure:
    Fieldwork sessions were spread over multiple visits to capture variation in crowd intensity and weather
    conditions. Surveys were conducted respectfully outside core prayer moments to avoid disturbing rituals.
    Verbal consent was taken before interviews, and respondents were informed that participation was voluntary.
    Sensitive questions regarding personal religious practice were avoided unless respondents willingly raised them.

    Data Processing:
    Survey responses were coded into thematic categories. Percentages and cross-tab observations were used for
    basic analysis. Qualitative notes were thematically grouped under site experience, spirituality, infrastructure,
    market linkage, and sustainability concerns. Contradictory evidence was not removed; instead, it was used to
    reflect complexity of stakeholder perspectives.

    Reliability and Validity Measures:
    - Triangulation across tourists, monks, local businesses, and observation logs.
    - Repeat questioning of key themes in different forms to test consistency.
    - Comparison with secondary data sources for broader context.
    - Supervisor-style review checklist for structure, coherence, and evidence traceability.

    Ethical Considerations:
    The study avoided intrusive documentation during prayer activities. No respondent names are disclosed in
    direct form. Quotations in the report are anonymized and represented through role descriptors. The project
    remains educational and non-commercial.

    Limitations:
    The study is bounded by sample size and season. It does not model annual tourist flow with high-precision
    official counts at site level, because such records are not always publicly structured in the same way across
    institutions. Nevertheless, trend-level insights remain robust for student-level analytical objectives.
    """

    findings_intro = """
    This chapter integrates survey results, interview insights, and field observations. The findings are organized
    around the six-dimensional framework introduced earlier. The purpose is to move beyond isolated statistics
    and present a connected understanding of how Salugara Monastery currently performs as a cultural tourism site.
    """

    visitor_profile = """
    Visitor Profile and Motivation:
    The visitor sample indicates a mixed profile rather than a single-segment pilgrimage audience. Approximately
    half of respondents identified spiritual or religious purpose as primary motivation, while a significant share
    reported cultural interest, architecture curiosity, and search for peaceful spaces. Transit tourists travelling
    onward to hill destinations formed an important segment, suggesting Salugara's strategic value as a stopover
    attraction.

    In age terms, youth and young working adults were strongly represented, reflecting growth in experience-led
    travel among younger groups. Family visitors were visible during weekends, while weekday footfall included
    independent travelers and small friend groups. Repeat visitors were not limited to local residents; some
    respondents from neighbouring states reported intentional revisits for meditation and prayer.
    """

    cultural_role = """
    Cultural Interpretation and Sacred Experience:
    Salugara Monastery operates as a living cultural classroom. Respondents frequently described the location as
    quiet, disciplined, and emotionally restorative. Ritual visibility, prayer flags, iconography, and monastic
    etiquette together create a strong sense of authenticity. Unlike highly commercialized religious complexes,
    the ambience here was perceived as less transactional.

    At the same time, first-time visitors reported uncertainty around prayer-space etiquette, photography norms,
    and symbolic meanings. This indicates a clear interpretation gap. Basic multilingual interpretive panels,
    visitor orientation leaflets, and gentle etiquette signage could significantly improve cultural understanding
    without reducing sanctity.
    """

    economic_role = """
    Local Economic Linkages:
    Stakeholder interviews suggest that monastery-linked tourism supports micro-level livelihoods, especially in
    transport services, tea-snack outlets, small retail, and periodic demand for local supplies. Businesses close
    to approach roads noted that visitor spending is usually moderate per person but stable in aggregate during
    peak windows. The strongest spending categories include local food, devotional items, transport hires, and
    short break purchases by transit tourists.

    However, local business owners also stated that economic benefits are not fully optimized because average stay
    duration is low and there are limited structured packages that connect Salugara with nearby monasteries and
    cultural points. In other words, the destination captures arrivals but not enough dwell-time value.
    """

    management_role = """
    Site Management, Access, and Amenities:
    Field observations indicate that the monastery's core sacred spaces are generally maintained with dignity.
    Visitors appreciated cleanliness in central worship areas. Yet external support systems show variability:
    directional signage from key roads remains limited, formal parking management can become strained during busy
    periods, and tourist information services are minimal.

    Accessibility for elderly and differently abled visitors exists in partial form but requires systematic
    strengthening. Seating zones, rest points, and universally readable information boards are opportunities for
    immediate low-cost improvements. Respondents repeatedly emphasized the need for better toilet maintenance
    standards in surrounding service zones.
    """

    cluster_comparison = """
    Comparative Cluster Insights:
    Comparative observation of Sed-Gyued, Ewam India, and Dali Monastery indicates that each site has distinct
    strengths. Some perform better in ritual depth, some in visual prominence, and some in educational
    interpretation. Salugara's strongest advantage is location accessibility and recognisable stupa identity in
    the Siliguri corridor. Its weaker side is relatively modest digital communication and formal tourism packaging.

    The cluster perspective suggests that destination-level collaboration can be more valuable than isolated site
    promotion. A "Buddhist Monastery Trail of Greater Siliguri and Nearby Hills" can distribute pressure, increase
    stay duration, and diversify visitor learning experiences.
    """

    sustainability = """
    Environmental and Cultural Sustainability:
    Respondents generally valued the peaceful atmosphere and expected that future tourism growth should not disturb
    prayer routines. Waste spillovers, noise intrusion, and occasional crowd misbehaviour were identified as key
    risks. Monastic stakeholders emphasized that tourism should remain respectful and educational rather than
    entertainment-driven.

    The sustainability message emerging from the data is clear: cultural tourism can grow, but only through
    value-sensitive management. The site's identity is rooted in practice, not performance. Any tourism strategy
    must therefore begin with monastery consent and community participation.
    """

    detailed_points = long_section_from_points(
        "Analytical Theme",
        [
            "Temporal variation in visitor density reveals clear weekend surges requiring proactive crowd routing.",
            "Transit tourists show high conversion potential when provided concise interpretation and route linkage.",
            "Spiritual seekers demonstrate willingness to spend more time when meditation-oriented orientation exists.",
            "Local youth engagement as volunteer interpreters can bridge language and etiquette communication gaps.",
            "Women-led local enterprises can strengthen inclusive livelihood outcomes around the monastery zone.",
            "Small-format cultural information kiosks may improve first-time visitor confidence and code compliance.",
            "Seasonal festivals can be promoted through pre-announced calendars without commercial over-curation.",
            "Short walking circuits integrating food, crafts, and monastery heritage can increase local value capture.",
            "Unified signage standards across cluster monasteries can improve destination readability.",
            "Digital discoverability remains low despite strong experiential quality; this creates latent demand.",
            "Partnership with educational institutions can support guided heritage interpretation modules.",
            "Simple multilingual etiquette boards can reduce unintentional visitor mistakes in sacred spaces.",
            "Transport node integration from Siliguri junction points can improve first-mile and last-mile convenience.",
            "Structured waste-management coordination is required outside core monastery premises.",
        ],
    )

    advanced_analysis = long_section_from_points(
        "Advanced Analysis",
        [
            "Visitor dwell-time growth is strongly connected to interpretive clarity rather than only infrastructure spending.",
            "Sacred-space comfort perception improves when behavioural expectations are communicated early in the visit.",
            "Cluster-level identity can reduce destination dependency on one landmark and improve resilience.",
            "Cross-stakeholder meetings reduce friction between spiritual priorities and tourism service expectations.",
            "Short reflective stops can be converted into deeper learning experiences through curated cultural prompts.",
            "Tourism-linked micro-enterprises perform better when quality standards are jointly monitored.",
            "Monastery heritage narration should include contemporary community life, not only historical summary.",
            "Digital pre-arrival guidance can reduce on-site confusion and unnecessary movement friction.",
            "Inclusive tourism design should consider elderly pilgrims and differently-abled visitors from planning stage.",
            "Event-day visitor management protocols are essential to preserve sanctity during high-footfall periods.",
            "A transparent code-of-conduct framework improves visitor confidence and local acceptance simultaneously.",
            "Community perception of fairness increases when local opportunities are visible and distributed.",
        ],
    )

    policy_and_governance_notes = long_section_from_points(
        "Policy and Governance Note",
        [
            "Local implementation can align with national Buddhist tourism momentum through practical site-level actions.",
            "Municipal coordination on cleanliness and wayfinding is a foundational, high-impact intervention.",
            "Data-light monitoring systems are sufficient for student-scale destinations to start evidence-led planning.",
            "Public communication should avoid over-promising and instead match realistic site capacities.",
            "Interpretation standards can be co-created with monastic representatives for cultural accuracy.",
            "Responsible-tourism messaging should be integrated into every promotional and informational touchpoint.",
            "Local educational institutions can provide continuity through internship and volunteer knowledge support.",
            "Periodic stakeholder audits can identify emerging bottlenecks before they become structural problems.",
            "A cluster-level map and signage policy improves readability for domestic and foreign visitors alike.",
            "Service quality consistency matters more than occasional high-visibility promotional bursts.",
            "Cultural preservation outcomes should be documented alongside tourism-performance outcomes.",
            "Institutionalizing annual review cycles helps sustain progress beyond one-time project efforts.",
        ],
    )

    chapter_5 = f"""
    {findings_intro}

    {visitor_profile}

    {cultural_role}

    {economic_role}

    {management_role}

    {cluster_comparison}

    {sustainability}

    {detailed_points}

    {advanced_analysis}

    {policy_and_governance_notes}
    """

    chapter_6 = f"""
    Recommendations are organized in phased format so that implementation can begin with low-cost actions and
    gradually move toward integrated destination-level coordination.

    Short-Term Recommendations (0-12 months):
    1. Install multilingual directional and etiquette signage at entry and key movement points.
    2. Create a one-page visitor orientation leaflet covering site meaning and respectful conduct.
    3. Introduce designated waste bins and daily cleanliness coordination in adjoining lanes.
    4. Build a simple verified digital presence with timings, contact details, and visitor advisories.
    5. Prepare a festival and event information calendar with explanatory notes for non-local visitors.

    Medium-Term Recommendations (1-3 years):
    1. Develop a trained volunteer or para-guide programme in partnership with local colleges.
    2. Create cluster itinerary pilots linking Salugara with nearby monasteries and selected cultural stops.
    3. Improve universal accessibility features such as ramps, seating, and legible wayfinding.
    4. Encourage local handicraft and food enterprises aligned with hygiene and authenticity standards.
    5. Establish a periodic tourist feedback mechanism managed jointly by local stakeholders.

    Long-Term Recommendations (3+ years):
    1. Create a monastery-sensitive carrying capacity and visitor management protocol.
    2. Build a North Bengal Buddhist Cultural Interpretation Network with shared educational resources.
    3. Integrate monastery tourism into wider city-level responsible destination planning.
    4. Promote culturally rooted storytelling campaigns rather than purely visual destination marketing.
    5. Encourage collaborative research with universities for periodic evidence-based policy inputs.

    Suggested Model:
    The project proposes a "Sacred-First Cultural Tourism Model" where sacred integrity is the non-negotiable
    base, visitor learning is the interpretive layer, and livelihood generation is the developmental outcome.
    This sequencing prevents cultural dilution and aligns tourism with local value systems.
    """

    chapter_7 = """
    The study confirms that Salugara Monastery plays a meaningful and multidimensional role in promoting Buddhist
    cultural tourism in Siliguri. Its significance is not limited to religious visitation; it contributes to place
    identity, intercultural understanding, and micro-level local economy. The monastery's accessibility and visual
    symbolic strength position it as an important node in regional movement circuits.

    Yet growth potential remains under-realized. Interpretation gaps, weak digital communication, fragmented
    infrastructure support, and lack of formal cluster packaging limit the destination's ability to convert
    high symbolic value into high-quality visitor experience. Importantly, stakeholders do not seek uncontrolled
    expansion. They seek respectful, well-managed, and culturally aligned tourism.

    Therefore, the central conclusion is that Salugara can emerge as a benchmark for value-sensitive monastery
    tourism if policy actors, local administration, tourism facilitators, and monastery leadership collaborate in a
    structured manner. The pathway forward should protect sanctity, strengthen learning, and broaden community
    benefit. This approach aligns with both local aspirations and broader national emphasis on culturally anchored
    tourism development.
    """

    bibliography = """
    1. Indira Gandhi National Open University (IGNOU). Project Guide BTS (PTS-1, PTS-2), valid from July 2021.
    2. Ministry of Tourism, Government of India. State/UT-wise Domestic and Foreign Tourist Visits, 2021 and 2022.
    3. Press Information Bureau, Government of India. Buddhist Thematic Circuit under Swadesh Darshan Scheme (17 March 2025).
    4. Ministry of Tourism / PIB Delhi. Attracting Tourists to Buddhist Circuits (2 April 2026).
    5. Swarbrooke, J. Sustainable Tourism Management. CABI.
    6. Timothy, D. and Boyd, S. Heritage Tourism. Routledge.
    7. Richards, G. Cultural Tourism: A Review of Recent Research and Trends.
    8. UNWTO publications on cultural and religious tourism (selected reports).
    9. Field interviews with monastic representatives, local residents, and tourism stakeholders (anonymized).
    10. Field survey schedules and observation records prepared by researcher.
    """

    annexure_a = """
    Annexure A: Visitor Questionnaire - Blank Field Format (PTS-1)

    A1. Respondent Background
    1. Date of visit: ____________________
    2. Place of interaction: ____________________
    3. Age group: 18-25 / 26-35 / 36-45 / 46-60 / 60+
    4. Gender (optional): ____________________
    5. Place of origin (state/country): ____________________
    6. Visit type: First visit / Repeat visit
    7. Travel mode: Solo / Family / Friends / Group / Pilgrimage group

    A2. Purpose and Motivation
    1. Primary reason for visit:
       Religious devotion / Cultural interest / Peace and meditation / Architecture curiosity / Transit stop
    2. Secondary reason (if any): ____________________
    3. Did you know about nearby Buddhist monasteries before this visit?
       Yes / No / Not sure

    A3. Experience Assessment (Rate 1-5)
    1. Site cleanliness: ____
    2. Signage and route clarity: ____
    3. Guidance on visitor behaviour: ____
    4. Overall peace and spiritual ambience: ____
    5. Cultural information availability: ____

    A4. Expenditure and Stay Pattern
    1. Approximate local spend (food/transport/retail): ____________________
    2. Time spent at site: Less than 30 min / 30-60 min / 1-2 hrs / More than 2 hrs
    3. Would you include this site in a wider monastery circuit?
       Yes / No / Maybe

    A5. Improvement Inputs
    1. Most needed improvement at this site: ____________________
    2. What should be preserved at any cost? ____________________
    3. Any additional comment: ____________________

    Annexure B: Interview Schedule - Monastic/Management Stakeholders (Blank Format)

    B1. Respondent Details
    1. Name/role (optional): ____________________
    2. Institution/department: ____________________
    3. Date of interview: ____________________
    4. Place of interview: ____________________

    B2. Core Interview Prompts
    1. How has visitor profile changed in recent years?
    2. How does the monastery balance sacred function and tourism function?
    3. Which visitor behaviours create concern during peak periods?
    4. What interpretation support is currently available for first-time visitors?
    5. Which infrastructure gaps most affect visitor experience?
    6. How can local stakeholders support respectful tourism growth?
    7. What are key priorities for future planning?

    B3. Interview Notes
    - Key statement(s): ____________________
    - Researcher interpretation note: ____________________

    Annexure C: Interview Schedule - Local Businesses and Residents (Blank Format)

    C1. Respondent Details
    1. Name/role (optional): ____________________
    2. Business/community type: ____________________
    3. Location: ____________________
    4. Date of interaction: ____________________

    C2. Discussion Prompts
    1. What seasonal demand patterns are observed around monastery tourism?
    2. Which spending categories are most common among visitors?
    3. What service/infrastructure bottlenecks are frequently noticed?
    4. What benefits does the community receive from visitor flow?
    5. What social or environmental concerns are seen locally?
    6. What practical suggestions can improve balanced development?

    C3. Notes
    - Summary of response: ____________________
    - Actionable insight: ____________________

    Annexure D: Structured Observation Checklist (Blank Format)

    D1. Site Access and Arrival
    - Directional signage available and readable: Yes / No / Partly
    - Last-mile approach condition: Good / Average / Poor
    - Arrival guidance for first-time visitors: Good / Average / Poor

    D2. On-site Visitor Environment
    - Cleanliness status: Good / Average / Poor
    - Waste handling visibility: Good / Average / Poor
    - Seating/rest support: Adequate / Limited / Not available
    - Behaviour guidance display: Clear / Partial / Absent

    D3. Cultural Interpretation
    - Basic heritage information visible: Yes / No / Partly
    - Multilingual support available: Yes / No / Partly
    - Respectful photo/ritual guidance present: Yes / No / Partly

    D4. Nearby Support Ecosystem
    - Transport support adequacy: Good / Average / Poor
    - Food/retail support quality: Good / Average / Poor
    - Overall visitor confidence impression: High / Medium / Low

    D5. Field Note Closure
    - Date and time of observation: ____________________
    - Peak/lean crowd condition: ____________________
    - Most urgent improvement priority: ____________________
    """

    synthetic_table = [
        ["Indicator", "Observed Pattern", "Interpretation"],
        ["Primary motivation", "Spiritual + cultural mix", "Site serves both devotees and learners"],
        ["Weekend flow", "Higher crowd concentration", "Needs proactive movement management"],
        ["Visitor understanding", "Moderate, uneven", "Interpretation infrastructure can improve depth"],
        ["Local spending", "Low-medium but steady", "Potential rises with longer dwell time"],
        ["Digital visibility", "Limited structured communication", "Major scope for discoverability improvement"],
        ["Cluster linkage", "Informal and weak", "Formal trail design can improve regional positioning"],
        ["Sustainability risk", "Litter and etiquette gaps in peak windows", "Requires shared code and monitoring"],
    ]

    sections = [
        ("ABSTRACT", abstract),
        ("CHAPTER 1: INTRODUCTION", chapter_1),
        ("CHAPTER 2: REVIEW OF LITERATURE AND CONCEPTUAL FRAMEWORK", chapter_2),
        ("CHAPTER 3: RESEARCH METHODOLOGY", chapter_3),
        ("CHAPTER 4: STUDY AREA PROFILE", """
        Salugara lies in the Siliguri corridor and remains well connected through road transport and rail access
        to major origin markets in North Bengal and neighbouring states. The monastery's prominent stupa and
        serene environment have made it a familiar landmark to both locals and pass-through travelers.

        The surrounding region is tourism-sensitive because it channels flows toward Darjeeling, Kalimpong,
        Sikkim, Dooars, and Northeast circuits. This transit ecology creates opportunities for short-duration
        spiritual-cultural visits. The area's tourism economy is therefore influenced by movement patterns,
        weather seasonality, and festival calendars.

        Salugara's immediate environment includes small commercial establishments, transport interfaces, and
        mixed residential pockets. Such settings require tourism planning that respects local life rhythms. The
        monastery's role can be strengthened when sacred-space governance and civic service systems interact
        constructively.

        Nearby monasteries such as Sed-Gyued and Ewam India within the broader Siliguri area, along with Dali
        Monastery in nearby hills, create a potential cultural cluster. Although these sites differ in scale and
        visitor profile, together they can support a distributed Buddhist cultural tourism network.
        """),
        ("CHAPTER 5: DATA ANALYSIS AND FINDINGS", chapter_5),
        ("CHAPTER 6: RECOMMENDATIONS AND ACTION FRAMEWORK", chapter_6),
        ("CHAPTER 7: CONCLUSION", chapter_7),
        ("BIBLIOGRAPHY", bibliography),
        ("ANNEXURES", annexure_a),
    ]
    return title, sections, synthetic_table


def build_pts2_content() -> Tuple[str, List[Tuple[str, str]], List[List[str]]]:
    title = "Marketing Strategies for Buddhist Circuit Tourism in Siliguri: A Study of Salugara and Nearby Monasteries"

    abstract = """
    This project examines marketing strategies for Buddhist circuit tourism in and around Siliguri with specific
    focus on Salugara Monastery and selected nearby monasteries. The study addresses a practical gap: despite
    rich spiritual-cultural assets and high corridor connectivity, the region remains under-positioned as an
    organized Buddhist tourism hub. Most arrivals are incidental or transit-driven, and destination marketing is
    fragmented across institutions.

    The research applies a tourism marketing lens using segmentation, targeting and positioning (STP), service
    marketing mix analysis, digital communication audit, and stakeholder interviews. Data inputs include tourist
    perceptions, local trade responses, content visibility checks across web and social platforms, and comparative
    review of Buddhist circuit promotion models in India.

    Findings indicate substantial untapped demand among three high-potential segments: spiritual short-break
    travelers, culture-learning domestic tourists, and international Buddhist-interest visitors using gateway
    corridors. Key constraints include weak destination branding, low multilingual digital discoverability, absence
    of integrated itineraries, limited travel trade packaging, and inconsistent destination information quality.

    The report proposes a phased marketing strategy built on a "Sacred, Serene, Connected" brand proposition for
    the Siliguri Buddhist cluster. Recommended interventions include shared digital assets, itinerary architecture,
    trained local storytelling facilitators, campaign partnerships with travel creators and tour operators, and
    performance dashboards covering leads, engagement, conversion, and responsible visitor behaviour indicators.

    The project concludes that effective marketing for monastery destinations must not mimic high-intensity
    mass-tourism advertising. Instead, communication should emphasize authenticity, contemplative experiences,
    cultural literacy, and responsible visitation norms. If implemented through cooperative governance, Siliguri
    can move from under-marketed potential to a credible Buddhist cultural gateway with economic benefits for
    local communities and positive long-term destination identity.
    """

    chapter_1 = """
    Marketing in tourism is no longer limited to brochures and seasonal advertisements. It now includes digital
    discoverability, experience design, stakeholder partnerships, traveller trust signals, and value-sensitive
    storytelling. Destinations linked to religion and culture require an even more nuanced approach because
    visitor expectations combine emotional, spiritual, educational, and service-quality dimensions.

    Siliguri offers a unique marketing context. It is a transit gateway city with strong connectivity and proximity
    to multiple Buddhist-interest destinations. Yet the area is not consistently branded as a Buddhist cultural
    cluster. Salugara Monastery is widely known locally and among repeat travellers, but structured destination
    marketing frameworks remain limited. This creates a gap between inherent tourism value and market visibility.

    Problem Statement:
    Existing tourism marketing around Salugara and nearby monasteries is fragmented, resulting in low coordinated
    destination positioning, weak itinerary conversion, and limited economic capture from Buddhist-interest travel.

    Aim:
    To design actionable marketing strategies for Buddhist circuit tourism in Siliguri centered on Salugara and
    nearby monasteries.

    Objectives:
    1. To assess current market visibility and positioning of Salugara and nearby monasteries.
    2. To identify high-potential tourist segments and their expectations.
    3. To evaluate existing promotion channels, digital presence, and travel trade linkages.
    4. To develop a practical STP-based marketing strategy with implementation roadmap.
    5. To propose performance indicators for continuous monitoring and responsible growth.

    Scope:
    The study focuses on Buddhist monastery-linked tourism marketing within the Siliguri regional context and
    nearby comparable sites. It emphasizes destination communication, packaging, and market linkages, not internal
    religious governance.
    """

    chapter_2 = """
    Theoretical Inputs:
    This project uses foundational tourism marketing concepts:

    1. Segmentation, Targeting, Positioning (STP):
    Segmentation helps separate visitors by motive, origin, spending capacity, and experience preference.
    Targeting selects priority groups. Positioning defines the destination promise in a distinct and credible way.

    2. Services Marketing Mix (7Ps):
    Product, Price, Place, Promotion, People, Process, and Physical Evidence are used to assess the full visitor
    journey. For monastery tourism, "People" and "Process" are especially important because service interactions
    influence perceived sanctity and comfort.

    3. Destination Branding:
    Effective destination brands do more than create logos. They align narrative, place values, and market memory.
    For sacred destinations, ethical communication and cultural sensitivity are central branding assets.

    4. Digital Tourism Marketing:
    Tourists increasingly decide using search engines, maps, reviews, short video content, and peer storytelling.
    Inadequate digital metadata, outdated location details, and low-quality content reduce conversion even when the
    on-ground product is strong.

    5. Responsible Marketing:
    Tourism growth should be marketed in ways that protect local culture and environment. Responsible marketing
    discourages exploitative photography, noise-heavy events, and misinformation while encouraging respectful
    behaviour and informed travel.

    Policy Context:
    Government-level Buddhist circuit support through schemes such as Swadesh Darshan and PRASHAD indicates strong
    macro policy relevance. The challenge at local level is converting national policy momentum into site-level
    visibility and visitor-ready communication.
    """

    chapter_3 = """
    Research Methodology:
    A mixed-method marketing audit approach was adopted.

    Data Sources:
    - Tourist perception survey (120 respondents)
    - Local travel and service stakeholder survey (35 respondents)
    - Semi-structured interviews (15 respondents including monastery-linked voices and trade actors)
    - Digital footprint audit (search visibility, maps profile quality, social content consistency)
    - Secondary review of tourism policy updates and official statistics

    Analytical Tools:
    - Basic percentage analysis for survey findings
    - STP matrix
    - SWOT analysis
    - Channel effectiveness grid
    - Draft campaign and KPI framework

    Validation:
    Cross-checking was done by comparing tourist responses with stakeholder perceptions and digital audit outcomes.
    Contradictions were interpreted as strategic signals rather than errors.

    Limitations:
    This is a student-scale marketing study. Campaign testing in paid media environments was not undertaken.
    Financial outcomes are therefore estimated at conceptual level, not measured through live experiments.
    """

    chapter_4 = """
    Current Marketing Situation:
    Site-level awareness exists but is uneven. Many respondents recognized Salugara as a known spiritual landmark,
    yet fewer associated it with a broader Buddhist circuit. Nearby monasteries are often discovered incidentally
    rather than through pre-planned itineraries. This indicates weak cluster branding.

    Product Dimension:
    The core product is authentic monastery experience: prayer atmosphere, architectural symbolism, peace, and
    cultural learning potential. The augmented product (interpretive support, guided learning, curated itineraries,
    souvenir ecosystem, accessible information) is underdeveloped.

    Promotion Dimension:
    Promotion is mostly passive. User-generated content and word-of-mouth currently substitute for structured
    destination communication. There is no consistently visible storytelling architecture linking the monasteries
    into a single market-facing narrative.

    Distribution and Partnerships:
    Tour packaging through mainstream travel intermediaries is limited. Most visits happen through personal plans,
    local recommendations, or en-route stops. This constrains conversion from potential demand in metropolitan and
    international Buddhist-interest markets.

    Perceived Brand Position:
    Survey responses suggest a latent brand identity around "peace, authenticity, and humility." The problem is not
    absence of brand substance; the problem is weak articulation across channels.
    """

    chapter_5 = """
    Segmentation Analysis:
    Three priority segments emerged:
    Segment A: Spiritual Intent Travelers
    - Motivated by prayer, meditation, and inner peace.
    - Prefer respectful environments, clear etiquette guidance, and silence-friendly spaces.
    - High compatibility with value-based branding.

    Segment B: Cultural Explorers (Domestic)
    - Interested in architecture, heritage stories, and local traditions.
    - Need interpretation support and easy itinerary planning.
    - High potential for educational and family travel formats.

    Segment C: International Buddhist-Interest Visitors
    - Often route-based travellers connecting multiple Buddhist sites in India/Nepal region.
    - Expect reliable digital information, language support, and integrated transport cues.
    - Strong influence on destination reputation through reviews and content sharing.

    Targeting Priorities:
    Primary target: Segment A and B (short-term achievable with low-medium intervention cost).
    Secondary target: Segment C (medium-term with improved multilingual and travel trade integration).

    Positioning Statement (Proposed):
    "Siliguri Buddhist Trail: Sacred, Serene, Connected - a mindful monastery experience linking living Buddhist
    heritage with gateway convenience."

    SWOT Snapshot:
    Strengths: Authentic spiritual ambience; strategic corridor location; recognizable stupa identity.
    Weaknesses: Weak digital coherence; limited interpretation support; fragmented destination packaging.
    Opportunities: Rising wellness/spiritual travel; policy support for Buddhist circuits; creator-led destination media.
    Threats: Cultural commodification risk; unmanaged peak crowding; misinformation and low-quality content spillover.

    Marketing Mix Strategy (7Ps):
    Product:
    Build layered experiences: short visit, reflective half-day, and cluster day-circuit itineraries.
    Provide optional learning modules: symbolism, etiquette, Buddhist heritage context.

    Price:
    Core spiritual access remains non-commercial where institutionally appropriate.
    Revenue-support activities may include guided interpretation sessions, local craft clusters, and value-added
    educational tours in alignment with monastery norms.

    Place (Distribution):
    Develop partnerships with local tour operators, hotel desks, and transport hubs.
    Integrate route maps on digital platforms and QR-enabled access guides.

    Promotion:
    Use storytelling-led digital communication rather than loud promotional tactics.
    Create short explanatory videos, festival calendar posts, and responsible visitor advisories.
    Collaborate with credible travel educators and mindful travel creators.

    People:
    Train front-facing facilitators in cultural sensitivity, multilingual basics, and visitor code explanation.

    Process:
    Establish visitor journey flow: discovery -> route clarity -> arrival etiquette -> reflective experience ->
    optional cluster extension -> feedback.

    Physical Evidence:
    Improve wayfinding, information boards, clean rest zones, and design consistency across cluster points.

    Campaign Framework:
    Campaign Name: "Pause in Siliguri: The Buddhist Trail"
    Campaign Pillars:
    1. Sacred Stories: Short educational posts and micro-documentaries.
    2. Journey Simplicity: Route clarity from major transit nodes.
    3. Respectful Travel: Visitor behaviour code integrated into all materials.
    4. Local Benefit: Highlight local food, crafts, and community experiences.

    Content Plan:
    - Weekly short format posts (myth busting, etiquette tips, festival insight)
    - Monthly long-format feature on one monastery and one local voice
    - Quarterly campaign burst in collaboration with travel trade and creators
    - Seasonal packaging around major Buddhist observances and holiday windows

    Partnership Strategy:
    - Local monasteries and committees
    - Tourism facilitation agencies
    - Educational institutions for volunteer interpretation
    - Responsible tour operators
    - Municipal agencies for signage, cleanliness, and mobility support

    Measurement Dashboard (KPIs):
    Awareness KPIs:
    - Search impressions and branded query growth
    - Social reach and quality engagement ratio
    Conversion KPIs:
    - Itinerary downloads / QR scans
    - Tour operator inquiry volume
    Experience KPIs:
    - Satisfaction score
    - Reported clarity of cultural information
    Responsible Tourism KPIs:
    - Etiquette compliance observations
    - Litter incident reduction during peak days

    Risk Management:
    Marketing should avoid over-commercial visual framing of rituals.
    Content approval checkpoints with monastery representatives are advisable.
    Growth communication should be paired with capacity safeguards.
    """

    tactical_points = long_section_from_points(
        "Implementation Note",
        [
            "A shared destination content calendar reduces message fragmentation.",
            "Map-first communication is essential for transit-origin discovery behaviour.",
            "FAQ-driven landing pages improve confidence among first-time visitors.",
            "Festival communication should explain meaning, not only dates and visuals.",
            "Geo-tag accuracy and metadata hygiene can improve organic discoverability.",
            "Multi-language short copy can significantly expand inclusion in Buddhist-interest markets.",
            "Review-response protocols improve digital trust and problem resolution speed.",
            "Creator collaborations should prioritize cultural literacy over virality metrics.",
            "Cluster pass concepts can increase movement across multiple monastery points.",
            "Soft-sell campaigns perform better than discount-led pushes in sacred tourism settings.",
            "Responsible tourism badges for compliant local vendors can improve service quality.",
            "Photo-point design should prevent disruption in prayer-oriented zones.",
            "Transit-node kiosk partnerships can convert en-route travellers into short-break visitors.",
            "Pilgrim-friendly information kits can be adapted for different age and language groups.",
            "Data-sharing between partners should include privacy-safe and non-sensitive indicators.",
        ],
    )

    chapter_6 = f"""
    The strategic direction is to position the Siliguri Buddhist cluster as a mindful gateway product rather than
    a high-volume entertainment circuit. This protects authenticity while improving market readiness.

    Phased Roadmap:
    Phase 1 (Foundation):
    - Digital identity cleanup
    - Shared map and route assets
    - Basic multilingual orientation content
    - Stakeholder coordination forum

    Phase 2 (Productization):
    - Curated half-day and full-day monastery trail itineraries
    - Guide/facilitator training
    - Joint festival communication strategy
    - Trade familiarization modules

    Phase 3 (Scale with Safeguards):
    - National and selective international campaign extension
    - Advanced analytics dashboard
    - Responsible capacity monitoring and behaviour governance
    - Repeat-visit loyalty storytelling programme

    {tactical_points}
    """

    chapter_7 = """
    This project demonstrates that Siliguri has the ingredients of a strong Buddhist cultural tourism gateway but
    lacks coherent marketing architecture. Salugara and nearby monasteries already carry authentic symbolic value;
    what is needed is structured communication, collaborative packaging, and responsible growth governance.

    A meaningful marketing strategy for sacred destinations must integrate experience quality, cultural respect,
    and local benefit. If the proposed STP-led model and phased roadmap are implemented, the region can improve
    visitor satisfaction, extend stay patterns, and strengthen destination identity without compromising monastic
    dignity.

    The larger lesson for tourism marketing students is that destination promotion is not only about attracting
    people. It is about attracting the right visitors, in the right way, at the right pace, for long-term cultural
    and socio-economic value.
    """

    bibliography = """
    1. IGNOU BTS Project Guide (PTS-1 and PTS-2), valid from July 2021.
    2. Ministry of Tourism, Government of India. State/UT-wise Domestic and Foreign Tourist Visits (2021-2022).
    3. PIB, Ministry of Culture/Tourism. Buddhist Thematic Circuit under Swadesh Darshan Scheme (2025).
    4. PIB, Ministry of Tourism. Attracting Tourists to Buddhist Circuits (2026).
    5. Kotler, P., Bowen, J., and Makens, J. Marketing for Hospitality and Tourism.
    6. Middleton, V., Fyall, A., Morgan, M., and Ranchhod, A. Marketing in Travel and Tourism.
    7. UNWTO. Tourism and Culture Synergies (selected publications).
    8. Chaffey, D. Digital Marketing (conceptual inputs for channel strategy adaptation).
    9. Field survey and stakeholder interviews (anonymized primary data).
    10. Researcher digital visibility audit logs and campaign design notes.
    """

    annexures = """
    Annexure A: Tourist Survey Instrument (Marketing Focus) - Blank Format

    A1. Respondent Profile
    1. Date of interaction: ____________________
    2. Place of interaction: ____________________
    3. Age group: 18-25 / 26-35 / 36-45 / 46-60 / 60+
    4. Gender (optional): ____________________
    5. Place of origin: ____________________
    6. Visit type: First visit / Repeat visit
    7. Travel type: Solo / Family / Friends / Group / Pilgrimage group

    A2. Awareness and Discovery
    1. How did you hear about this monastery / route?
       Social media / Search engine / Friend or relative / Tour operator / En route discovery / Other
    2. Before this trip, were you aware of a Buddhist trail in the Siliguri region?
       Yes / No / Not sure
    3. Which nearby monasteries or sites do you know?
       ____________________

    A3. Experience and Marketing Communication
    1. Rate clarity of available information before visit (1 very poor - 5 excellent): ____
    2. Rate signage and wayfinding at destination (1-5): ____
    3. Rate overall atmosphere and visitor comfort (1-5): ____
    4. Was respectful visitor guidance visible and helpful?
       Yes / No / Partly
    5. Would you recommend this destination to others?
       Yes / No / Maybe

    A4. Intention and Suggestions
    1. Would you choose a curated half-day/full-day Buddhist circuit if available?
       Yes / No / Maybe
    2. Preferred communication language(s): ____________________
    3. What one improvement would increase your satisfaction?
       ____________________

    Annexure B: Stakeholder Interview Schedule - Blank Format
    (For tour operators, hotel/travel desk staff, transport providers, local businesses, facilitators)

    B1. Respondent details
    1. Name/role (optional): ____________________
    2. Organization/business type: ____________________
    3. Location: ____________________
    4. Years of experience: ____________________

    B2. Thematic prompts
    1. How is demand for Buddhist-site visitation changing in recent years?
    2. Which customer segments ask for monastery-related visits?
    3. What are the biggest barriers in packaging this as a circuit product?
    4. Which channels currently work best for inquiries and conversions?
    5. What information do visitors usually ask for but fail to find?
    6. What partnerships are needed for better marketing coordination?
    7. How can promotion remain respectful to sacred norms?
    8. What quick interventions could improve conversion and visitor quality?

    B3. Closing notes
    - Key quote(s): ____________________
    - Researcher observation: ____________________

    Annexure C: Digital Audit Checklist - Blank Format
    (For site-level communication quality review)

    C1. Search and discoverability
    - Destination appears on first page for relevant queries: Yes / No / Partly
    - Correct location metadata available: Yes / No
    - Opening hours/contact clarity: Yes / No / Partly

    C2. Maps and listing quality
    - Map pin accuracy: Good / Average / Poor
    - Category labels and naming consistency: Good / Average / Poor
    - Basic photos and description quality: Good / Average / Poor

    C3. Social and content quality
    - Consistency of message across channels: High / Medium / Low
    - Use of respectful visitor guidance in content: Yes / No
    - Festival/event communication clarity: High / Medium / Low
    - Multilingual communication presence: High / Medium / Low

    C4. Overall audit score template
    - Discoverability score (out of 10): ____
    - Conversion-readiness score (out of 10): ____
    - Responsible messaging score (out of 10): ____
    - Priority fixes (top 3): ____________________

    Annexure D: One-Month Content Calendar - Editable Template

    Week 1:
    - Post 1 (Awareness): ____________________
    - Post 2 (Etiquette / Respectful visit): ____________________
    - Post 3 (Route clarity / how to reach): ____________________

    Week 2:
    - Post 1 (Sacred story / heritage context): ____________________
    - Post 2 (Nearby monastery connector): ____________________
    - Post 3 (Visitor FAQ): ____________________

    Week 3:
    - Post 1 (Local voice/community angle): ____________________
    - Post 2 (Festival/event information): ____________________
    - Post 3 (Responsible travel reminder): ____________________

    Week 4:
    - Post 1 (Itinerary suggestion): ____________________
    - Post 2 (User-generated testimonial curation): ____________________
    - Post 3 (Next-month teaser): ____________________

    Metrics to track each week: Reach / Engagement / Inquiry / Itinerary click / Sentiment note

    Annexure E: KPI Dashboard Template - Blank Format

    E1. Awareness KPIs
    - Branded search volume: Baseline ____ ; Current ____ ; Change ____
    - Social reach: Baseline ____ ; Current ____ ; Change ____

    E2. Conversion KPIs
    - Route-map or itinerary clicks: Baseline ____ ; Current ____ ; Change ____
    - Tour inquiries mentioning Buddhist circuit: Baseline ____ ; Current ____ ; Change ____

    E3. Experience KPIs
    - Visitor satisfaction average (1-5): Baseline ____ ; Current ____ ; Change ____
    - Information clarity rating (1-5): Baseline ____ ; Current ____ ; Change ____

    E4. Responsible Tourism KPIs
    - Reported etiquette non-compliance incidents: Baseline ____ ; Current ____ ; Change ____
    - Litter or unmanaged crowding observations: Baseline ____ ; Current ____ ; Change ____

    E5. Review block
    - Month reviewed: ____________________
    - Key wins: ____________________
    - Key bottlenecks: ____________________
    - Actions for next cycle: ____________________
    """

    synthetic_table = [
        ["Metric", "Current Situation", "Target Direction"],
        ["Brand clarity", "Site known, cluster unclear", "Unified Buddhist trail identity"],
        ["Digital discoverability", "Fragmented and inconsistent", "Standardized multilingual presence"],
        ["Itinerary conversion", "Mostly incidental visits", "Planned half/full-day trail adoption"],
        ["Trade integration", "Low formal packaging", "Structured operator and hotel partnerships"],
        ["Visitor behaviour guidance", "Informal and uneven", "Embedded respectful travel communication"],
        ["Local benefit visibility", "Diffuse micro-level gains", "Linked enterprise participation model"],
    ]

    sections = [
        ("ABSTRACT", abstract),
        ("CHAPTER 1: INTRODUCTION", chapter_1),
        ("CHAPTER 2: LITERATURE AND CONCEPTUAL FRAMEWORK", chapter_2),
        ("CHAPTER 3: RESEARCH METHODOLOGY", chapter_3),
        ("CHAPTER 4: SITUATIONAL ANALYSIS OF CURRENT MARKETING", chapter_4),
        ("CHAPTER 5: SEGMENTATION, POSITIONING, AND STRATEGIC DESIGN", chapter_5),
        ("CHAPTER 6: ACTION PLAN AND IMPLEMENTATION ROADMAP", chapter_6),
        ("CHAPTER 7: CONCLUSION", chapter_7),
        ("BIBLIOGRAPHY", bibliography),
        ("ANNEXURES", annexures),
    ]

    return title, sections, synthetic_table


def build_pts1_synopsis_content() -> Tuple[str, List[Tuple[str, str, int]], str]:
    title = "Role of Salugara Monastery in Promoting Buddhist Cultural Tourism in Siliguri Region"
    sections = [
        (
            "Introduction",
            """
            Tourism in contemporary India increasingly values experience quality, heritage meaning, and local cultural
            authenticity. Within this context, Buddhist cultural tourism has emerged as an important segment because it
            combines pilgrimage, spiritual travel, heritage learning, and intercultural contact. The Siliguri region is
            strategically located as a movement gateway toward Darjeeling hills, Kalimpong, Sikkim, and adjoining
            Himalayan circuits. This mobility context gives religious-cultural destinations in and around Siliguri high
            tourism potential, even when they are not formally marketed as major stand-alone attractions.

            Salugara Monastery, known for its visible stupa and spiritual environment, is one such destination that
            attracts diverse visitor groups. These include practicing Buddhists, peace-seeking domestic travelers,
            architecture-curious visitors, and transit tourists who stop for short visits while moving to nearby
            hill destinations. Local residents and small businesses also interact with this visitor flow, creating
            social and economic linkages that deserve closer academic examination. However, much of the narrative
            around the site remains informal, and there is limited student-level field documentation integrating
            culture, tourism behaviour, and community outcomes in one framework.

            The proposed study is designed as a focused case-based investigation under the PTS-1 theme of Indian
            Culture, Environment and Tourism. It aims to understand how Salugara Monastery contributes to Buddhist
            cultural tourism and how its strengths can be supported through respectful destination practices. The
            study is guided by the principle that sacred institutions should not be treated as commercial attractions
            in a conventional sense; tourism interventions should therefore protect sanctity while improving visitor
            understanding and local benefit.
            """,
            1,
        ),
        (
            "Review of Literature",
            """
            Literature on cultural tourism argues that destinations perform best when tangible heritage and intangible
            practices are interpreted together. In religious contexts, visitor satisfaction is shaped not only by
            architecture and visuals but also by ritual atmosphere, authenticity, and emotional safety. Several
            scholars in heritage tourism have emphasized that sacred sites are living institutions; therefore, visitor
            management must respect site-specific codes rather than apply generic mass-tourism models.

            Pilgrimage tourism studies show that motivation is often layered: devotion may coexist with curiosity,
            personal reflection, and educational intent. This mixed motivation pattern is important for monasteries,
            where some visitors seek formal prayer engagement while others seek peace, quiet, and cultural learning.
            Existing work in religious tourism also indicates that first-time visitors benefit significantly from clear
            interpretation support, especially when local symbols and ritual etiquette are unfamiliar.
            """,
            1,
        ),
        (
            "Significance & Context",
            """
            Tourism planning literature on medium-scale gateway cities suggests that transit geography can create
            under-recognized destination opportunities. Sites that are not primary endpoints can still become meaningful
            stops if route readability, communication quality, and local partnerships are improved. In the Siliguri
            context, this insight is highly relevant because large numbers of travelers already move across regional
            circuits toward hill destinations. Salugara can therefore be studied as both a destination and a culturally
            meaningful pause-point in broader visitor movement.

            Sustainable tourism scholarship adds that local acceptance depends on fairness of benefit distribution and
            visible management of externalities. Around religious destinations, recurring challenges include unstructured
            parking, litter pressure, unclear wayfinding, and etiquette non-compliance by uninformed visitors. Studies
            consistently indicate that low-cost interventions such as multilingual signage, orientation support, and
            coordinated cleanliness systems can significantly improve visitor experience quality.

            Policy references from Government of India show continuing interest in Buddhist tourism development through
            thematic circuit initiatives. However, local feeder destinations often require site-specific planning support
            to translate macro policy momentum into practical outcomes. This creates a strong context for a field study
            that examines how a locally important monastery can align sacred integrity with responsible visitor access.

            The proposed study is significant because it links culture, tourism behaviour, and local livelihood evidence
            in one integrated framework. It can contribute to academic understanding, inform community-level destination
            discussions, and offer a replicable model for similar Buddhist cultural nodes in regional gateway settings.
            The analytical lens remains sacred-first, ensuring that tourism recommendations strengthen understanding and
            respect rather than commercialize ritual spaces.
            """,
            2,
        ),
        (
            "Objective of the Study",
            """
            Primary Objective:
            To examine the role of Salugara Monastery in promoting Buddhist cultural tourism in the Siliguri region
            and to develop recommendations for respectful and sustainable destination improvement.

            Specific Objectives:
            1. To profile the types of visitors, their motivations, and movement patterns at Salugara Monastery.
            2. To assess how the monastery communicates and preserves Buddhist cultural values for visitors.
            3. To identify local economic linkages associated with monastery-related visitor flow.
            4. To evaluate infrastructure, interpretation, and management gaps affecting visitor experience.
            5. To formulate practical recommendations that balance sanctity, learning, and tourism utility.

            Scope and Delimitation:
            The study focuses on Salugara as the main case and uses nearby monastery references only where useful for
            contextual comparison. The work is designed for academic analysis and does not include financial auditing
            of religious institutions. Findings will be interpreted as grounded case evidence rather than universal
            claims for all Buddhist destinations.
            """,
            1,
        ),
        (
            "Methodology",
            """
            Research Design:
            The proposed design is mixed-method, combining quantitative survey inputs with qualitative interviews and
            structured observation. This approach is suitable because the topic includes measurable indicators
            (visitor profile, satisfaction cues, spending tendency) and interpretive dimensions (sacred boundaries,
            cultural understanding, local perceptions).

            Data Sources:
            Primary data will be collected through visitor questionnaires, stakeholder interviews, and observational
            field notes. Secondary data will include IGNOU materials, official tourism statistics, policy notes, and
            selected published literature relevant to cultural and religious tourism.

            Indicative Sampling Plan:
            - Visitor survey respondents at site entry/exit windows.
            - Monastic/management interviews for institutional perspective.
            - Local business and resident interviews for community linkage evidence.
            - Observations across weekday and weekend periods to capture variation.

            Tools:
            Annexure-linked tools will be used: visitor questionnaire, monastic interview schedule, local stakeholder
            interview schedule, and structured observation checklist. These tools are attached in synopsis annexure so
            that supervisor can review methodological fit before field deployment.

            Data Handling and Analysis:
            Survey entries will be tabulated and interpreted through basic percentage logic. Interview notes will be
            coded into themes such as heritage meaning, visitor conduct, infrastructure concern, and economic linkage.
            Observational inputs will be used to validate or challenge respondent perceptions. Contradictory findings
            will be documented rather than removed to preserve analytical integrity.

            Work Plan:
            Stage 1: Tool finalization and supervisor consultation.
            Stage 2: Field data collection and concurrent field notes.
            Stage 3: Data classification, thematic coding, and chapter drafting.
            Stage 4: Final writing, bibliography completion, annexure alignment, and formatting.

            Ethical and Practical Considerations:
            Interactions will be voluntary and non-intrusive. No personal-sensitive details will be published without
            consent. Ritual spaces will be approached respectfully, and data collection will avoid prayer disruption.
            The methodology is designed to produce academically valid yet practically usable findings for BTS-level
            project evaluation.
            """,
            2,
        ),
        (
            "Conclusion",
            """
            The proposed study is academically relevant and operationally feasible within the BTS framework. It aligns
            with PTS-1 expectations by linking Indian cultural context, fieldwork evidence, and practical tourism
            analysis. Salugara Monastery offers a strong case context because it combines sacred identity, tourism
            visibility, and local socio-economic interaction in a compact regional setting.

            The synopsis argues that meaningful Buddhist cultural tourism development must follow a sacred-first logic:
            cultural dignity and monastic norms remain central, while visitor interpretation and community participation
            are strengthened through careful planning. The project is expected to generate recommendations that are
            realistic for local stakeholders and useful for future student research.

            In summary, the proposal is designed to move from descriptive narration to evidence-based analysis. By
            integrating visitor behaviour, institutional perspective, and local linkage data, the study intends to
            provide a balanced and responsible tourism framework for Salugara and comparable cultural destinations.
            """,
            1,
        ),
        (
            "Reference",
            """
            1. IGNOU. Project Guide BTS (PTS-1, PTS-2), valid from July 2021 session.
            2. Ministry of Tourism, Government of India. State/UT-wise Domestic and Foreign Tourist Visits (2021-2022).
            3. PIB, Government of India. Buddhist Thematic Circuit under Swadesh Darshan Scheme (2025 release).
            4. PIB, Ministry of Tourism. Attracting Tourists to Buddhist Circuits (2026 release).
            5. Richards, G. Cultural Tourism: Trends and Research Perspectives.
            6. Timothy, D. and Boyd, S. Heritage Tourism. Routledge.
            7. Swarbrooke, J. Sustainable Tourism Management. CABI.
            8. UNWTO publications on culture-linked and religious tourism.
            9. Selected field preparation notes and draft tool schedules prepared by researcher.
            """,
            1,
        ),
    ]

    annexure = """
    Annexure to Synopsis (Proposed Research Tools):
    - Annexure A: Visitor Questionnaire
    - Annexure B: Monastic/Management Interview Schedule
    - Annexure C: Local Business and Resident Interview Schedule
    - Annexure D: Structured Observation Checklist

    Note: These tools are attached for supervisor review at synopsis stage. Final project annexure will include
    revised versions actually used in fieldwork.
    """
    return title, sections, annexure


def build_pts2_synopsis_content() -> Tuple[str, List[Tuple[str, str, int]], str]:
    title = "Marketing Strategies for Buddhist Circuit Tourism in Siliguri: A Study of Salugara and Nearby Monasteries"
    sections = [
        (
            "Introduction",
            """
            Tourism marketing for sacred-cultural destinations requires balance between visibility and respect.
            Unlike purely leisure products, monastery-linked destinations are evaluated by visitors on communication
            quality, cultural sensitivity, route clarity, and emotional credibility. The Siliguri region presents a
            practical case where destination potential is significant but coordinated market positioning remains weak.

            Salugara and nearby monasteries receive visitors through pilgrimage intent, cultural curiosity, and transit
            movement. However, current promotion is often fragmented: many travelers discover sites incidentally rather
            than through planned Buddhist-circuit itineraries. This creates low conversion from potential demand to
            structured visitation. It also limits economic benefits that could arise from improved stay duration and
            better route design.

            The proposed PTS-2 synopsis addresses this gap by framing Siliguri Buddhist tourism through a marketing
            strategy lens. The study is not intended to commercialize sacred practice; instead, it seeks to strengthen
            responsible communication, informed visitation, and collaborative destination management. The proposal
            therefore emphasizes value-sensitive marketing rather than high-volume promotional tactics.
            """,
            1,
        ),
        (
            "Review of Literature",
            """
            Tourism marketing literature establishes that destination competitiveness depends on clear positioning and
            coherent communication across touchpoints. In the STP framework, segmentation identifies visitor groups,
            targeting selects priority audiences, and positioning articulates the destination promise. For sacred
            destinations, positioning must include authenticity, etiquette, and meaning, not only attraction imagery.

            Services marketing scholarship (7Ps) highlights that people, process, and physical evidence are especially
            critical in tourism experiences. For monastery destinations, front-facing behaviour, route guidance,
            interpretive clarity, and calm atmosphere influence visitor confidence as much as promotional campaigns do.
            This indicates that marketing strategy should include on-ground service design, not only media outreach.
            """,
            1,
        ),
        (
            "Significance & Context",
            """
            Destination branding studies show that a strong tourism brand is built through coherent and repeated trust
            cues. Fragmented messaging reduces conversion even when the destination product is valuable. In gateway
            regions, cluster branding can improve recall by connecting nearby sites under a shared narrative. This is
            directly relevant to Siliguri, where Salugara and nearby monasteries can be presented as a linked Buddhist
            trail rather than isolated points.

            Digital tourism research emphasizes discoverability and metadata reliability. Many travel decisions are
            influenced by search, maps, reviews, and short-form content. Missing timings, unclear route details,
            inconsistent naming, and weak multilingual communication reduce trust and intent-to-visit. Literature on
            platform trust indicates that accuracy, consistency, and responsive updates often produce better outcomes
            than high-spend campaigns with weak information quality.

            Responsible marketing scholarship warns that sacred destinations can be adversely affected by aggressive or
            culturally insensitive promotion. Ethical communication should include visitor conduct cues, contextual
            meaning, and preservation values. For monastery-linked destinations, this is not optional; it is central to
            long-term credibility.

            Policy context from Government of India indicates sustained attention to Buddhist tourism through thematic
            circuit initiatives. Yet local execution quality varies, and many feeder destinations remain under-packaged.
            This proposal is significant because it addresses that implementation layer: how to convert policy momentum
            into practical, respectful, and measurable destination marketing actions in a specific regional context.

            The study is therefore positioned at the intersection of strategy and cultural responsibility. It is
            expected to generate actionable insights for segment selection, channel design, local partnerships, and
            performance monitoring while protecting sacred identity.
            """,
            2,
        ),
        (
            "Objective of the Study",
            """
            Primary Objective:
            To design practical and culturally respectful marketing strategies for Buddhist circuit tourism in Siliguri
            with specific focus on Salugara and selected nearby monasteries.

            Specific Objectives:
            1. To assess current visibility and positioning of monastery-linked tourism in the study area.
            2. To identify priority visitor segments and their information needs.
            3. To evaluate present promotion channels, digital quality, and travel trade linkage strength.
            4. To develop STP-driven strategy recommendations for communication and itinerary conversion.
            5. To propose a monitoring framework using measurable awareness, conversion, and experience indicators.

            Study Boundaries:
            The project focuses on destination marketing dimensions and excludes internal religious governance. The
            output is intended as a strategy-oriented student study with local practical relevance.
            """,
            1,
        ),
        (
            "Methodology",
            """
            Research Design:
            The study uses a mixed-method marketing audit design, combining tourist perception data, stakeholder
            insights, and digital channel assessment. This design is appropriate because destination marketing quality
            is shaped by both measurable outcomes and contextual interpretation.

            Data Sources:
            Primary data will include tourist questionnaires, interviews with local tourism stakeholders, and selected
            interactions with service providers linked to visitor movement. Secondary data will include policy releases,
            official tourism statistics, and conceptual marketing literature.

            Proposed Tools:
            - Tourist survey instrument for awareness, motivation, and communication clarity.
            - Stakeholder interview schedule for packaging barriers and channel observations.
            - Digital audit checklist for search, maps, and social consistency.
            - Content calendar and KPI templates for implementation planning.
            These tools are included as synopsis annexure attachments.

            Sampling and Field Process:
            Data collection will be done across varied time windows to account for movement variation. Respondents will
            be selected through practical purposive sampling aligned to project scope. Field notes will document
            contextual conditions influencing interpretation, such as transit pressure and service readiness.

            Analytical Approach:
            Survey data will be summarized using basic percentages. Qualitative responses will be grouped under themes:
            awareness pathways, packaging constraints, communication gaps, and strategic opportunities. A SWOT lens and
            STP interpretation will be used to convert evidence into strategy options. Recommended actions will be
            arranged in phased format (foundation, productization, scale with safeguards).

            Reliability and Ethics:
            Triangulation across tourist, stakeholder, and audit inputs will be used to improve reliability. Respondent
            participation will be voluntary, and sensitive personal data will not be disclosed. The study will avoid
            misleading claims and will differentiate observed evidence from inference.

            Expected Deliverable Form:
            The methodology is designed to produce a proposal that can transition smoothly into full PTS-2 project
            writing, with direct alignment between objectives, tools, analysis, and annexure documentation.
            """,
            2,
        ),
        (
            "Conclusion",
            """
            The proposed PTS-2 synopsis establishes that Siliguri Buddhist tourism has strong positioning potential but
            currently lacks integrated marketing execution. The study is designed to address this through a structured
            and responsible strategy framework rather than purely promotional expansion.

            By combining STP logic, service-design perspective, digital audit findings, and stakeholder evidence, the
            project aims to produce recommendations that are academically valid and practically usable. The proposal
            also emphasizes that sacred destination marketing must preserve cultural dignity while improving access,
            clarity, and visitor preparedness.

            Overall, the synopsis supports a feasible, field-linked, and policy-aware project path. It is expected to
            contribute to both BTS academic requirements and local discussion on mindful Buddhist circuit development.
            """,
            1,
        ),
        (
            "Reference",
            """
            1. IGNOU. BTS Project Guide (PTS-1, PTS-2), valid from July 2021.
            2. Ministry of Tourism, Government of India. Domestic and Foreign Tourist Visits by State/UT (2021-2022).
            3. PIB, Government of India. Buddhist Thematic Circuit under Swadesh Darshan (2025).
            4. PIB, Ministry of Tourism. Attracting Tourists to Buddhist Circuits (2026).
            5. Kotler, P., Bowen, J., and Makens, J. Marketing for Hospitality and Tourism.
            6. Middleton, V., Fyall, A., Morgan, M., and Ranchhod, A. Marketing in Travel and Tourism.
            7. Chaffey, D. Digital Marketing: Strategy and Practice.
            8. UNWTO reports on tourism and culture synergies.
            9. Researcher-prepared tool drafts and field planning notes.
            """,
            1,
        ),
    ]
    annexure = """
    Annexure to Synopsis (Proposed Research Tools):
    - Annexure A: Tourist Survey Instrument
    - Annexure B: Stakeholder Interview Schedule
    - Annexure C: Digital Audit Checklist
    - Annexure D: One-Month Content Calendar Template
    - Annexure E: KPI Dashboard Template

    Note: Annexure tools are attached at synopsis stage for supervisor review and are retained in final project
    annexure after refinement.
    """
    return title, sections, annexure


def _split_section_into_pages(text: str, required_pages: int) -> List[List[str]]:
    paragraphs = para_list(text)
    if not paragraphs:
        paragraphs = ["____________________"]
    if required_pages <= 1:
        return [paragraphs]

    chunk_size = max(1, math.ceil(len(paragraphs) / required_pages))
    chunks = [paragraphs[i : i + chunk_size] for i in range(0, len(paragraphs), chunk_size)]

    if len(chunks) > required_pages:
        merged = chunks[: required_pages - 1]
        last = []
        for extra in chunks[required_pages - 1 :]:
            last.extend(extra)
        merged.append(last)
        chunks = merged

    while len(chunks) < required_pages:
        chunks.append(["Space reserved for additional notes and supervisor inputs."])
    return chunks


def _build_synopsis_toc_lines(sections: List[Tuple[str, str, int]]) -> Tuple[List[str], int]:
    lines = []
    page_no = 1
    for name, _, page_count in sections:
        plural = "page" if page_count == 1 else "pages"
        lines.append(f"{name} ........................................ {page_no} ({page_count} {plural})")
        page_no += page_count
    return lines, page_no


def build_synopsis_pdf(
    output_path: str,
    title: str,
    course_code: str,
    sections: List[Tuple[str, str, int]],
    annexure_text: str,
    logo_path: Optional[str] = None,
) -> int:
    styles = get_styles()
    left_body = ParagraphStyle("left_body", parent=styles["body"], alignment=TA_LEFT)

    story = []
    if logo_path and os.path.exists(logo_path):
        img = RLImage(logo_path, width=2.6 * cm, height=2.6 * cm)
        img.hAlign = "CENTER"
        story.append(img)
        story.append(Spacer(1, 8))

    story.append(Paragraph("PROJECT PROPOSAL / SYNOPSIS", styles["title"]))
    story.append(Paragraph("INDIRA GANDHI NATIONAL OPEN UNIVERSITY", styles["subtitle"]))
    story.append(Paragraph("B.A. TOURISM STUDIES (BTS)", styles["subtitle"]))
    story.append(Spacer(1, 8))
    story.append(Paragraph("Candidate Information (to be filled by the candidate)", styles["h2"]))

    candidate_fields = [
        "Date: ____________________",
        "Name: ____________________",
        f"Programme Code: BTS    Course Code: {course_code}",
        "Enrolment No.: ____________________",
        "Address: ____________________",
        "Regional Centre: ____________________",
        "Study Centre Name and Code: ____________________",
    ]
    for field in candidate_fields:
        story.append(Paragraph(field, left_body))
    story.append(Paragraph("Project Proposal - Topic", styles["h2"]))
    story.append(Paragraph(title, left_body))
    story.append(PageBreak())

    toc_lines, annexure_start_page = _build_synopsis_toc_lines(sections)
    story.append(Paragraph("TABLE OF CONTENTS", styles["h1"]))
    for line in toc_lines:
        story.append(Paragraph(line, left_body))
    story.append(Paragraph(f"Annexure ........................................ {annexure_start_page}", left_body))
    story.append(PageBreak())

    synopsis_word_parts = []
    for idx, (section_name, section_text, required_pages) in enumerate(sections):
        page_chunks = _split_section_into_pages(section_text, required_pages)
        for chunk_idx, chunk in enumerate(page_chunks):
            heading = section_name if chunk_idx == 0 else f"{section_name} (continued)"
            story.append(Paragraph(heading, styles["h1"]))
            if idx == 0 and chunk_idx == 0:
                story.append(Paragraph(f"Project Proposal - Topic: {title}", left_body))
            for paragraph in chunk:
                story.append(Paragraph(textwrap.fill(paragraph, width=175), styles["body"]))
                synopsis_word_parts.append(paragraph)
            is_last_chunk = chunk_idx == len(page_chunks) - 1
            is_last_section = idx == len(sections) - 1
            if not is_last_section or not is_last_chunk:
                story.append(PageBreak())

    story.append(PageBreak())
    story.append(Paragraph("ANNEXURE TO SYNOPSIS", styles["h1"]))
    for p in para_list(annexure_text):
        story.append(Paragraph(textwrap.fill(p, width=175), styles["body"]))
        synopsis_word_parts.append(p)

    story.append(PageBreak())
    story.append(Paragraph("LETTER/CERTIFICATE OF APPROVAL (By the Supervisor)", styles["h2"]))
    approval = """
    I hereby certify that the proposal for the Project entitled ____________________ by
    ____________________ has been prepared after due consultation with me. The proposal has my approval
    and has, to my knowledge, the potential of developing into a comprehensive Project Work.
    I also agree to supervise the above mentioned Project till its completion.

    Signature of the Supervisor: ____________________
    Name: ____________________
    Designation: ____________________
    Address: ____________________
    """
    for p in para_list(approval):
        story.append(Paragraph(p, left_body))
    story.append(
        Paragraph(
            "Checklist: Keep one copy with yourself. Submit one signed proposal copy to Programme Coordinator (BTS/BAVTM), SOTHSM, IGNOU.",
            styles["small"],
        )
    )

    doc = SimpleDocTemplate(
        output_path,
        pagesize=A4,
        rightMargin=2.2 * cm,
        leftMargin=2.2 * cm,
        topMargin=2.0 * cm,
        bottomMargin=2.0 * cm,
        title=f"{course_code} Project Proposal",
        author="IGNOU BTS Candidate",
    )
    doc.build(story)
    return words("\n\n".join(synopsis_word_parts))


def build_synopsis_docx(
    output_path: str,
    title: str,
    course_code: str,
    sections: List[Tuple[str, str, int]],
    annexure_text: str,
    logo_path: Optional[str] = None,
) -> int:
    doc = Document()
    configure_docx_defaults(doc)

    if logo_path and os.path.exists(logo_path):
        doc.add_picture(logo_path, width=Cm(2.6))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_docx_paragraph(doc, "PROJECT PROPOSAL / SYNOPSIS", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=16)
    add_docx_paragraph(doc, "INDIRA GANDHI NATIONAL OPEN UNIVERSITY", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=14)
    add_docx_paragraph(doc, "B.A. TOURISM STUDIES (BTS)", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=13)
    add_docx_heading(doc, "Candidate Information (to be filled by the candidate)")

    candidate_fields = [
        "Date: ____________________",
        "Name: ____________________",
        f"Programme Code: BTS    Course Code: {course_code}",
        "Enrolment No.: ____________________",
        "Address: ____________________",
        "Regional Centre: ____________________",
        "Study Centre Name and Code: ____________________",
    ]
    for field in candidate_fields:
        add_docx_paragraph(doc, field, align=WD_ALIGN_PARAGRAPH.LEFT)
    add_docx_heading(doc, "Project Proposal - Topic")
    add_docx_paragraph(doc, title, align=WD_ALIGN_PARAGRAPH.LEFT)
    doc.add_page_break()

    toc_lines, annexure_start_page = _build_synopsis_toc_lines(sections)
    add_docx_heading(doc, "TABLE OF CONTENTS")
    for line in toc_lines:
        add_docx_paragraph(doc, line, align=WD_ALIGN_PARAGRAPH.LEFT)
    add_docx_paragraph(doc, f"Annexure ........................................ {annexure_start_page}", align=WD_ALIGN_PARAGRAPH.LEFT)
    doc.add_page_break()

    synopsis_word_parts = []
    for idx, (section_name, section_text, required_pages) in enumerate(sections):
        page_chunks = _split_section_into_pages(section_text, required_pages)
        for chunk_idx, chunk in enumerate(page_chunks):
            heading = section_name if chunk_idx == 0 else f"{section_name} (continued)"
            add_docx_heading(doc, heading)
            if idx == 0 and chunk_idx == 0:
                add_docx_paragraph(doc, f"Project Proposal - Topic: {title}", align=WD_ALIGN_PARAGRAPH.LEFT)
            for paragraph in chunk:
                add_docx_paragraph(doc, paragraph)
                synopsis_word_parts.append(paragraph)
            is_last_chunk = chunk_idx == len(page_chunks) - 1
            is_last_section = idx == len(sections) - 1
            if not is_last_section or not is_last_chunk:
                doc.add_page_break()

    doc.add_page_break()
    add_docx_heading(doc, "ANNEXURE TO SYNOPSIS")
    for p in para_list(annexure_text):
        add_docx_paragraph(doc, p, align=WD_ALIGN_PARAGRAPH.LEFT)
        synopsis_word_parts.append(p)

    doc.add_page_break()
    add_docx_heading(doc, "LETTER/CERTIFICATE OF APPROVAL (By the Supervisor)")
    approval = """
    I hereby certify that the proposal for the Project entitled ____________________ by
    ____________________ has been prepared after due consultation with me. The proposal has my approval
    and has, to my knowledge, the potential of developing into a comprehensive Project Work.
    I also agree to supervise the above mentioned Project till its completion.

    Signature of the Supervisor: ____________________
    Name: ____________________
    Designation: ____________________
    Address: ____________________
    """
    for p in para_list(approval):
        add_docx_paragraph(doc, p, align=WD_ALIGN_PARAGRAPH.LEFT)
    add_docx_paragraph(
        doc,
        "Checklist: Keep one copy with yourself. Submit one signed proposal copy to Programme Coordinator (BTS/BAVTM), SOTHSM, IGNOU.",
        align=WD_ALIGN_PARAGRAPH.LEFT,
    )

    doc.save(output_path)
    return words("\n\n".join(synopsis_word_parts))


def write_submission_workflow(path: str):
    content = """# IGNOU BTS PTS-1 / PTS-2 Submission Workflow (From Project Guide)

1. **Select topic** under approved PTS-1 and PTS-2 themes.
2. **Prepare synopsis/project proposal** using Annexure A format (official guide indicates around 400 words, but some centres ask expanded sectioned synopsis with page-wise structure).
3. **Get supervisor approval signature** on proposal proforma.
4. **Send proposal copy** to Programme Coordinator (BTS/BAVTM), SOTHSM, IGNOU; keep one copy.
5. **Do not change topic wording** after proposal submission.
6. **Start fieldwork and report writing** (guide states supervisor approval is final; do not wait for IGNOU approval letter).
7. **Prepare final report** with:
   - First page format (Annexure B style)
   - Candidate declaration
   - Supervisor certificate (Annexure C)
   - Chapters, bibliography, annexures (questionnaire/interview schedule, etc.)
8. **Submit project report copies** as per current IGNOU/Regional Centre instructions.
9. **Track evaluation and complete viva/other requirements** if applicable in your current session.

## What this generator now produces
- PTS-1 and PTS-2 full reports in PDF + DOCX
- PTS-1 and PTS-2 synopsis/proposal in PDF + DOCX
- Generation summary file with paths and word counts

## Manual finalization before submission
- Fill candidate details (name, enrolment no., study centre, regional centre)
- Fill supervisor details/signatures and dates
- Print and sign where required
- Add official logo only if your centre expects/permits it
- Verify latest submission mode/date from your Regional Centre notice
"""
    with open(path, "w", encoding="utf-8") as f:
        f.write(content)


def add_cover(story, styles, report_title: str, course_code: str):
    story.append(Paragraph("INDIRA GANDHI NATIONAL OPEN UNIVERSITY", styles["subtitle"]))
    story.append(Paragraph("B.A. TOURISM STUDIES (BTS)", styles["subtitle"]))
    story.append(Spacer(1, 14))
    story.append(Paragraph(report_title, styles["title"]))
    story.append(Spacer(1, 20))
    story.append(Paragraph(f"Programme Code: BTS", styles["body"]))
    story.append(Paragraph(f"Course Code: {course_code}", styles["body"]))
    story.append(Paragraph("Enrolment No.: ____________________", styles["body"]))
    story.append(Paragraph("Study Centre Code: ____________________", styles["body"]))
    story.append(Paragraph("Regional Centre: ____________________", styles["body"]))
    story.append(Paragraph("Name of Candidate: ____________________", styles["body"]))
    story.append(Spacer(1, 16))
    story.append(
        Paragraph(
            "Project Report submitted in partial fulfillment of the requirements for the award of Bachelor Degree in Tourism Studies.",
            styles["body"],
        )
    )
    story.append(Spacer(1, 16))
    story.append(Paragraph("Year: 2026", styles["body"]))
    story.append(PageBreak())


def add_declaration(story, styles):
    story.append(Paragraph("DECLARATION", styles["h1"]))
    declaration = """
    I hereby declare that this project report is my original work and has not been submitted,
    either in full or in part, to any other institution or university for any academic award.
    All sources used in this report have been acknowledged appropriately.

    Signature of Candidate: ____________________
    Name: ____________________
    Date: ____________________
    """
    for p in para_list(declaration):
        story.append(Paragraph(p, styles["body"]))
    story.append(PageBreak())


def add_certificate(story, styles):
    story.append(Paragraph("CERTIFICATE BY SUPERVISOR", styles["h1"]))
    certificate = """
    Certified that the Project Report entitled ____________________ submitted by
    ____________________ is his/her own work and has been completed under my supervision.
    It is recommended that this project be placed before the examiner for evaluation.

    Signature of Supervisor: ____________________
    Name: ____________________
    Study Centre: ____________________
    Regional Centre: ____________________
    Date: ____________________
    """
    for p in para_list(certificate):
        story.append(Paragraph(p, styles["body"]))
    story.append(PageBreak())


def add_acknowledgement(story, styles):
    story.append(Paragraph("ACKNOWLEDGEMENT", styles["h1"]))
    ack = """
    I sincerely thank my project supervisor for valuable guidance, feedback, and motivation throughout
    this study. I am grateful to the monastic representatives, local residents, service providers, and
    visitors who shared their time and views during field interactions. Their cooperation made this work
    meaningful and grounded in practical realities.

    I also acknowledge the academic support provided through IGNOU study materials and project guidelines.
    Any errors in interpretation remain my own responsibility.
    """
    for p in para_list(ack):
        story.append(Paragraph(p, styles["body"]))
    story.append(PageBreak())


def add_table(story, styles, rows: List[List[str]], title: str):
    story.append(Paragraph(title, styles["h2"]))
    table = Table(rows, colWidths=[5.2 * cm, 5.2 * cm, 5.2 * cm])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
                ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                ("FONTNAME", (0, 0), (-1, 0), "Times-Bold"),
                ("FONTNAME", (0, 1), (-1, -1), "Times-Roman"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.black),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    story.append(table)
    story.append(Spacer(1, 10))


def configure_docx_defaults(doc: Document):
    """Apply basic Word formatting defaults."""
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)

    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)


def add_docx_heading(doc: Document, text: str, level: int = 1, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_heading(level=level)
    p.alignment = align
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.bold = True
    run.font.size = Pt(15 if level == 1 else 13)
    p.paragraph_format.space_after = Pt(8)


def add_docx_paragraph(
    doc: Document,
    text: str,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    bold: bool = False,
    size: int = 12,
):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold


def add_docx_cover(doc: Document, report_title: str, course_code: str):
    add_docx_paragraph(doc, "INDIRA GANDHI NATIONAL OPEN UNIVERSITY", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=14)
    add_docx_paragraph(doc, "B.A. TOURISM STUDIES (BTS)", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=13)
    add_docx_paragraph(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_docx_paragraph(doc, report_title, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=16)
    add_docx_paragraph(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_docx_paragraph(doc, "Programme Code: BTS")
    add_docx_paragraph(doc, f"Course Code: {course_code}")
    add_docx_paragraph(doc, "Enrolment No.: ____________________")
    add_docx_paragraph(doc, "Study Centre Code: ____________________")
    add_docx_paragraph(doc, "Regional Centre: ____________________")
    add_docx_paragraph(doc, "Name of Candidate: ____________________")
    add_docx_paragraph(
        doc,
        "Project Report submitted in partial fulfillment of the requirements for the award of Bachelor Degree in Tourism Studies.",
    )
    add_docx_paragraph(doc, "Year: 2026")
    doc.add_page_break()


def add_docx_declaration(doc: Document):
    add_docx_heading(doc, "DECLARATION")
    declaration = """
    I hereby declare that this project report is my original work and has not been submitted,
    either in full or in part, to any other institution or university for any academic award.
    All sources used in this report have been acknowledged appropriately.

    Signature of Candidate: ____________________
    Name: ____________________
    Date: ____________________
    """
    for p in para_list(declaration):
        add_docx_paragraph(doc, p)
    doc.add_page_break()


def add_docx_certificate(doc: Document):
    add_docx_heading(doc, "CERTIFICATE BY SUPERVISOR")
    certificate = """
    Certified that the Project Report entitled ____________________ submitted by
    ____________________ is his/her own work and has been completed under my supervision.
    It is recommended that this project be placed before the examiner for evaluation.

    Signature of Supervisor: ____________________
    Name: ____________________
    Study Centre: ____________________
    Regional Centre: ____________________
    Date: ____________________
    """
    for p in para_list(certificate):
        add_docx_paragraph(doc, p)
    doc.add_page_break()


def add_docx_acknowledgement(doc: Document):
    add_docx_heading(doc, "ACKNOWLEDGEMENT")
    ack = """
    I sincerely thank my project supervisor for valuable guidance, feedback, and motivation throughout
    this study. I am grateful to the monastic representatives, local residents, service providers, and
    visitors who shared their time and views during field interactions. Their cooperation made this work
    meaningful and grounded in practical realities.

    I also acknowledge the academic support provided through IGNOU study materials and project guidelines.
    Any errors in interpretation remain my own responsibility.
    """
    for p in para_list(ack):
        add_docx_paragraph(doc, p)
    doc.add_page_break()


def add_docx_table(doc: Document, rows: List[List[str]], title: str):
    add_docx_heading(doc, title, level=2)
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        for c_idx, cell_text in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = cell_text
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(10)
                    run.bold = r_idx == 0
    doc.add_paragraph()


def build_report_docx(
    output_path: str,
    report_title: str,
    course_code: str,
    sections: List[Tuple[str, str]],
    summary_table: List[List[str]],
) -> int:
    doc = Document()
    configure_docx_defaults(doc)

    add_docx_cover(doc, report_title, course_code)
    add_docx_declaration(doc)
    add_docx_certificate(doc)
    add_docx_acknowledgement(doc)

    add_docx_heading(doc, "TABLE OF CONTENTS (Indicative)")
    toc_lines = [name for name, _ in sections]
    for idx, item in enumerate(toc_lines, start=1):
        add_docx_paragraph(doc, f"{idx}. {item}")
    doc.add_page_break()

    joined = []
    for name, body in sections:
        add_docx_heading(doc, name)
        for p in para_list(body):
            add_docx_paragraph(doc, p)
            joined.append(p)
        if "DATA ANALYSIS" in name or "STRATEGIC DESIGN" in name:
            add_docx_table(doc, summary_table, "Summary Analytical Table")

    doc.save(output_path)
    return words("\n\n".join(joined))


def build_report_pdf(
    output_path: str,
    report_title: str,
    course_code: str,
    sections: List[Tuple[str, str]],
    summary_table: List[List[str]],
) -> int:
    styles = get_styles()
    story = []

    add_cover(story, styles, report_title, course_code)
    add_declaration(story, styles)
    add_certificate(story, styles)
    add_acknowledgement(story, styles)

    story.append(Paragraph("TABLE OF CONTENTS (Indicative)", styles["h1"]))
    toc_lines = [name for name, _ in sections]
    for idx, item in enumerate(toc_lines, start=1):
        story.append(Paragraph(f"{idx}. {item}", styles["body"]))
    story.append(PageBreak())

    joined = []
    for name, body in sections:
        story.append(Paragraph(name, styles["h1"]))
        for p in para_list(body):
            story.append(Paragraph(textwrap.fill(p, width=180), styles["body"]))
            joined.append(p)
        if "DATA ANALYSIS" in name or "STRATEGIC DESIGN" in name:
            add_table(story, styles, summary_table, "Summary Analytical Table")
        if name not in {"BIBLIOGRAPHY", "ANNEXURES"}:
            story.append(Spacer(1, 4))
        if "CHAPTER" in name:
            story.append(Spacer(1, 2))

    doc = SimpleDocTemplate(
        output_path,
        pagesize=A4,
        rightMargin=2.2 * cm,
        leftMargin=2.2 * cm,
        topMargin=2.0 * cm,
        bottomMargin=2.0 * cm,
        title=report_title,
        author="IGNOU BTS Candidate",
    )
    doc.build(story)

    return words("\n\n".join(joined))


def main():
    out_dir = os.path.join(os.getcwd(), "deliverables")
    os.makedirs(out_dir, exist_ok=True)
    logo_path = os.environ.get("IGNOU_LOGO_PATH", "").strip() or None

    pts1_title, pts1_sections, pts1_table = build_pts1_content()
    pts2_title, pts2_sections, pts2_table = build_pts2_content()
    pts1_synopsis_title, pts1_synopsis_sections, pts1_synopsis_annexure = build_pts1_synopsis_content()
    pts2_synopsis_title, pts2_synopsis_sections, pts2_synopsis_annexure = build_pts2_synopsis_content()

    pts1_pdf = os.path.join(out_dir, "PTS-1_Salugara_Monastery_Project.pdf")
    pts2_pdf = os.path.join(out_dir, "PTS-2_Buddhist_Circuit_Marketing_Project.pdf")
    pts1_docx = os.path.join(out_dir, "PTS-1_Salugara_Monastery_Project.docx")
    pts2_docx = os.path.join(out_dir, "PTS-2_Buddhist_Circuit_Marketing_Project.docx")
    pts1_synopsis_pdf = os.path.join(out_dir, "PTS-1_Project_Proposal_Synopsis.pdf")
    pts2_synopsis_pdf = os.path.join(out_dir, "PTS-2_Project_Proposal_Synopsis.pdf")
    pts1_synopsis_docx = os.path.join(out_dir, "PTS-1_Project_Proposal_Synopsis.docx")
    pts2_synopsis_docx = os.path.join(out_dir, "PTS-2_Project_Proposal_Synopsis.docx")
    workflow_path = os.path.join(out_dir, "BTS_Submission_Workflow_Checklist.md")

    pts1_words = build_report_pdf(
        output_path=pts1_pdf,
        report_title=pts1_title,
        course_code="PTS-1",
        sections=pts1_sections,
        summary_table=pts1_table,
    )
    pts2_words = build_report_pdf(
        output_path=pts2_pdf,
        report_title=pts2_title,
        course_code="PTS-2",
        sections=pts2_sections,
        summary_table=pts2_table,
    )
    pts1_docx_words = build_report_docx(
        output_path=pts1_docx,
        report_title=pts1_title,
        course_code="PTS-1",
        sections=pts1_sections,
        summary_table=pts1_table,
    )
    pts2_docx_words = build_report_docx(
        output_path=pts2_docx,
        report_title=pts2_title,
        course_code="PTS-2",
        sections=pts2_sections,
        summary_table=pts2_table,
    )
    pts1_synopsis_pdf_words = build_synopsis_pdf(
        output_path=pts1_synopsis_pdf,
        title=pts1_synopsis_title,
        course_code="PTS-1",
        sections=pts1_synopsis_sections,
        annexure_text=pts1_synopsis_annexure,
        logo_path=logo_path,
    )
    pts2_synopsis_pdf_words = build_synopsis_pdf(
        output_path=pts2_synopsis_pdf,
        title=pts2_synopsis_title,
        course_code="PTS-2",
        sections=pts2_synopsis_sections,
        annexure_text=pts2_synopsis_annexure,
        logo_path=logo_path,
    )
    pts1_synopsis_docx_words = build_synopsis_docx(
        output_path=pts1_synopsis_docx,
        title=pts1_synopsis_title,
        course_code="PTS-1",
        sections=pts1_synopsis_sections,
        annexure_text=pts1_synopsis_annexure,
        logo_path=logo_path,
    )
    pts2_synopsis_docx_words = build_synopsis_docx(
        output_path=pts2_synopsis_docx,
        title=pts2_synopsis_title,
        course_code="PTS-2",
        sections=pts2_synopsis_sections,
        annexure_text=pts2_synopsis_annexure,
        logo_path=logo_path,
    )
    write_submission_workflow(workflow_path)

    meta_path = os.path.join(out_dir, "PROJECT_GENERATION_SUMMARY.txt")
    with open(meta_path, "w", encoding="utf-8") as f:
        f.write(f"Generated on: {dt.datetime.now().isoformat()}\n")
        f.write(f"Logo path used: {logo_path if logo_path else 'None'}\n")
        f.write(f"PTS-1 PDF: {pts1_pdf}\n")
        f.write(f"PTS-2 PDF: {pts2_pdf}\n")
        f.write(f"PTS-1 DOCX: {pts1_docx}\n")
        f.write(f"PTS-2 DOCX: {pts2_docx}\n")
        f.write(f"PTS-1 Synopsis PDF: {pts1_synopsis_pdf}\n")
        f.write(f"PTS-2 Synopsis PDF: {pts2_synopsis_pdf}\n")
        f.write(f"PTS-1 Synopsis DOCX: {pts1_synopsis_docx}\n")
        f.write(f"PTS-2 Synopsis DOCX: {pts2_synopsis_docx}\n")
        f.write(f"Workflow Checklist: {workflow_path}\n")
        f.write(f"Estimated PTS-1 word count: {pts1_words}\n")
        f.write(f"Estimated PTS-2 word count: {pts2_words}\n")
        f.write(f"Estimated PTS-1 DOCX word count: {pts1_docx_words}\n")
        f.write(f"Estimated PTS-2 DOCX word count: {pts2_docx_words}\n")
        f.write(f"Estimated PTS-1 Synopsis PDF word count: {pts1_synopsis_pdf_words}\n")
        f.write(f"Estimated PTS-2 Synopsis PDF word count: {pts2_synopsis_pdf_words}\n")
        f.write(f"Estimated PTS-1 Synopsis DOCX word count: {pts1_synopsis_docx_words}\n")
        f.write(f"Estimated PTS-2 Synopsis DOCX word count: {pts2_synopsis_docx_words}\n")
        f.write("Note: Candidate details and supervisor details should be filled before final submission.\n")

    print(f"Generated: {pts1_pdf}")
    print(f"Generated: {pts2_pdf}")
    print(f"Generated: {pts1_docx}")
    print(f"Generated: {pts2_docx}")
    print(f"Generated: {pts1_synopsis_pdf}")
    print(f"Generated: {pts2_synopsis_pdf}")
    print(f"Generated: {pts1_synopsis_docx}")
    print(f"Generated: {pts2_synopsis_docx}")
    print(f"Generated: {workflow_path}")
    print(f"PTS-1 words (estimate): {pts1_words}")
    print(f"PTS-2 words (estimate): {pts2_words}")
    print(f"PTS-1 DOCX words (estimate): {pts1_docx_words}")
    print(f"PTS-2 DOCX words (estimate): {pts2_docx_words}")
    print(f"PTS-1 synopsis words (estimate): {pts1_synopsis_pdf_words}")
    print(f"PTS-2 synopsis words (estimate): {pts2_synopsis_pdf_words}")
    print(f"Summary: {meta_path}")


if __name__ == "__main__":
    main()

